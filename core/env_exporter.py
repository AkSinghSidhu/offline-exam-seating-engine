"""
ENV exporter for generating Docx envelopes based on a template.
"""
from pathlib import Path
from typing import Dict, List, Tuple
from collections import defaultdict
from copy import deepcopy

from docx import Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

class EnvExporter:
    """Exports envelope pages based on the 13-5-26.docx template."""

    def export(self, plan, output_path: str, template_path: str, date_str: str, progress_callback=None) -> bool:
        try:
            doc = Document(template_path)
            
            # Extract the 10-element block which represents one full envelope page
            template_elements = [deepcopy(e) for e in doc.element.body[:10]]
            
            # Extract sectPr to preserve document validity
            sectPr = None
            for e in doc.element.body:
                if e.tag.endswith('sectPr'):
                    sectPr = deepcopy(e)
                    break
            
            # Clear the existing document body
            doc.element.body.clear()

            class_groups_by_room = self._group_students_by_room_and_class_ordered(plan)

            if not class_groups_by_room:
                return False

            total_groups = sum(len(groups) for _, groups in class_groups_by_room)
            processed = 0

            for room_plan, groups in class_groups_by_room:
                room_name = room_plan.room.name
                for class_key, students_info in sorted(groups.items()):
                    # Add a block for this class
                    self._add_envelope_block(doc, template_elements, room_name, class_key, students_info, date_str)
                    
                    processed += 1
                    if progress_callback:
                        progress_callback(int(processed / total_groups * 100))

            # Restore sectPr at the very end of the document body
            if sectPr is not None:
                doc.element.body.append(sectPr)

            output = Path(output_path)
            output.parent.mkdir(parents=True, exist_ok=True)
            doc.save(output)
            return True

        except Exception as e:
            print(f"ENV export error: {e}")
            import traceback
            traceback.print_exc()
            return False

    def _group_students_by_room_and_class_ordered(self, plan) -> List[Tuple[object, Dict[str, List[dict]]]]:
        results = []
        for room_plan in plan.room_plans:
            num_rows = len(room_plan.grid)
            num_cols = len(room_plan.grid[0]) if num_rows > 0 else 0
            
            groups: Dict[str, List[dict]] = defaultdict(list)

            for col in range(num_cols):
                for row in range(num_rows):
                    seat = room_plan.grid[row][col]
                    if seat and not seat.is_empty and seat.student:
                        student = seat.student
                        class_part = student.student_class or "Unknown"
                        sem_part = student.semester or ""
                        class_key = f"{class_part} Sem {sem_part}" if sem_part else class_part

                        groups[class_key].append({
                            'roll_number': student.roll_number,
                            'student_class': student.student_class,
                            'semester': student.semester,
                            'subject': student.subject,
                            'exam_id': seat.exam_id,
                            'course_code': getattr(student, 'course_code', ''),
                        })

            if groups:
                for key in groups:
                    groups[key].sort(key=lambda s: s['roll_number'])
                results.append((room_plan, dict(groups)))

        return results

    def _add_envelope_block(self, doc, template_elements, room_name: str, class_key: str, students_info: List[dict], date_str: str):
        # 1. Append deepcopies to the doc
        for e in template_elements:
            doc.element.body.append(deepcopy(e))
            
        table = doc.tables[-1]
        
        p_date = None
        p_room = None
        
        # Scan last few paragraphs to find DATE and ROOM NO
        for p in reversed(doc.paragraphs[-10:]):
            text = p.text.strip().upper()
            if text.startswith("DATE:"):
                p_date = p
            elif text.startswith("ROOM NO:"):
                p_room = p
                
        if p_date:
            orig_text = p_date.text
            leading_spaces = orig_text[:len(orig_text) - len(orig_text.lstrip())]
            self._replace_in_paragraph(p_date, f"{leading_spaces}DATE: {date_str}")
            
        if p_room:
            orig_text = p_room.text
            leading_spaces = orig_text[:len(orig_text) - len(orig_text.lstrip())]
            self._replace_in_paragraph(p_room, f"{leading_spaces}ROOM NO: {room_name}")
            
        # Fill table
        if len(table.rows) >= 2:
            row = table.rows[1]
            info = students_info[0]
            program = info.get('student_class', '')
            sem = info.get('semester', '')
            course_title = info.get('subject', '')
            course_code = info.get('course_code', '')
            num_students = len(students_info)
            num_papers = num_students + 1
            
            self._set_cell_text(row.cells[0], program)
            self._set_cell_text(row.cells[1], sem)
            self._set_cell_text(row.cells[2], course_title)
            self._set_cell_text(row.cells[3], course_code)
            self._set_cell_text(row.cells[4], str(num_students))
            self._set_cell_text(row.cells[5], str(num_papers))

    def _replace_in_paragraph(self, p, new_text: str):
        if not p.runs:
            p.add_run(new_text)
            return
            
        p.runs[0].text = new_text
        for r in p.runs[1:]:
            r.text = ""

    def _set_cell_text(self, cell, text):
        if cell.paragraphs:
            p = cell.paragraphs[0]
            if not p.runs:
                p.add_run(text)
            else:
                p.runs[0].text = text
                for r in p.runs[1:]:
                    r.text = ""
        else:
            cell.text = text
